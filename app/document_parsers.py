import csv
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree


TEXT_FACT_KEYWORDS = ('项目名称', '艺人', '城市', '场馆', '档期', '预计人数', '平均票价', '费用', '预算', '审批', '授权')
RISK_KEYWORDS = ('风险', '违约', '责任', '不确定', '缺失')
GATE_KEYWORDS = ('审批', '授权', '批文', '确认', '付款')
NUMBER_FIELDS = {
    'expected_attendance',
    'avg_ticket_price',
    'artist_fee',
    'venue_cost',
    'marketing_cost',
    'production_cost',
}
FIELD_ALIASES = {
    'name': {'项目名称', '项目', '名称', 'project', 'projectname'},
    'type': {'类型', '项目类型', 'type'},
    'artist_name': {'艺人', 'artist', '嘉宾', '艺人名称'},
    'city': {'城市', '站点', '地区', 'city'},
    'venue': {'场馆', 'venue', '场地'},
    'schedule': {'日期', '档期', '演出日期', 'schedule', 'date'},
    'expected_attendance': {'预计人数', '容量', '上座', '可售座位', 'expectedattendance', 'attendance'},
    'avg_ticket_price': {'平均票价', '票价', '客单价', 'avgticketprice', 'ticketprice'},
    'artist_fee': {'艺人费', '出场费', 'artistfee'},
    'venue_cost': {'场租', '场地成本', '场地费', 'venuecost'},
    'marketing_cost': {'宣发费', '投放预算', 'marketingcost'},
    'production_cost': {'制作费', '舞美搭建', 'productioncost'},
}


def resolve_local_file_path(evidence):
    meta = evidence.meta or {}
    for value in (meta.get('local_path'), meta.get('path'), evidence.file_url):
        if not value:
            continue
        candidate = str(value)
        if candidate.startswith('file://'):
            candidate = candidate[7:]
        path = Path(candidate).expanduser()
        if path.exists() and path.is_file():
            return path
    return None


def parse_document_evidence(evidence, file_kind: str):
    path = resolve_local_file_path(evidence)
    if not path:
        return _unavailable_document_result(evidence, '文件不在本地可读路径，保留证据并等待人工录入或下载签名接入。')

    try:
        if file_kind == 'docx':
            extracted_text = _extract_docx_text(path)
        elif file_kind == 'pdf':
            extracted_text = _extract_pdf_text(path)
        else:
            extracted_text = path.read_text(encoding='utf-8', errors='ignore')
    except Exception as exc:
        return _unavailable_document_result(evidence, f'资料解析失败：{exc}')

    return build_single_project_candidates(evidence, extracted_text)


def parse_spreadsheet_evidence(evidence):
    path = resolve_local_file_path(evidence)
    if not path:
        candidates = (evidence.meta or {}).get('candidate_projects') or []
        return {
            'file_kind': 'spreadsheet',
            'parse_scope': 'batch_projects',
            'candidate_projects': candidates,
            'sheets': [],
            'field_mapping': {},
            'requires_mapping_confirmation': True,
            'requires_human_verification': True,
            'note': '未找到本地可读文件，已回退读取 metadata.candidate_projects。',
        }

    try:
        sheets = _extract_spreadsheet_sheets(path)
    except Exception as exc:
        return {
            'file_kind': 'spreadsheet',
            'parse_scope': 'batch_projects',
            'candidate_projects': [],
            'sheets': [],
            'field_mapping': {},
            'requires_mapping_confirmation': True,
            'requires_human_verification': True,
            'error': f'表格解析失败：{exc}',
        }

    candidates, field_mapping = build_candidate_projects_from_sheets(sheets)
    return {
        'file_kind': 'spreadsheet',
        'parse_scope': 'batch_projects',
        'candidate_projects': candidates,
        'sheets': sheets,
        'field_mapping': field_mapping,
        'requires_mapping_confirmation': True,
        'requires_human_verification': True,
        'note': '候选项目已由表格表头映射生成，确认后才会创建正式项目。',
    }


def build_single_project_candidates(evidence, extracted_text: str):
    lines = [line.strip() for line in extracted_text.splitlines() if line.strip()]
    fact_lines = [line for line in lines if any(keyword in line for keyword in TEXT_FACT_KEYWORDS)]
    risk_lines = [line for line in lines if any(keyword in line for keyword in RISK_KEYWORDS)]
    gate_lines = [line for line in lines if any(keyword in line for keyword in GATE_KEYWORDS)]
    fact_content = '\n'.join(fact_lines or lines[:12])[:1200] or '资料已解析，但未抽取到明确正文。'

    candidate_risks = [
        {
            'title': _line_title(line, '资料解析风险'),
            'level': 'medium',
            'mitigation': '由负责人复核对应条款、数值和履约条件后再进入最终判断。',
        }
        for line in risk_lines[:5]
    ] or [{
        'title': '资料解析未核验风险',
        'level': 'medium',
        'mitigation': '由负责人确认解析字段、合同条款和关键数值后再进入最终判断。',
    }]

    candidate_gates = [
        {
            'name': _line_title(line, '资料人工核验'),
            'status': 'pending',
            'required_evidence': evidence.name,
            'owner_group': 'B',
        }
        for line in gate_lines[:5]
    ] or [{
        'name': '资料人工核验',
        'status': 'pending',
        'required_evidence': evidence.name,
        'owner_group': 'B',
    }]

    return {
        'file_kind': 'docx' if str(evidence.name).lower().endswith('.docx') else 'pdf',
        'parse_scope': 'single_project',
        'extracted_text': extracted_text[:4000],
        'candidate_facts': [{
            'title': f'{evidence.name} 解析候选事实',
            'content': fact_content,
            'source': evidence.source or evidence.evidence_type,
        }],
        'candidate_assumptions': [{
            'title': '资料解析结果待人工核验',
            'content': '解析出的项目字段、日期、费用和条款需由负责人复核后才能作为最终事实使用。',
            'confidence': 60 if fact_lines else 40,
        }],
        'candidate_risks': candidate_risks,
        'candidate_gates': candidate_gates,
        'requires_human_verification': True,
        'note': '解析结果已写入候选事实、假设、风险和门禁，状态保持待核验。',
    }


def build_candidate_projects_from_sheets(sheets):
    all_candidates = []
    merged_mapping = {}
    for sheet in sheets:
        rows = sheet.get('rows') or []
        if not rows:
            continue
        headers = [str(value).strip() if value is not None else '' for value in rows[0]]
        mapping = {header: _field_for_header(header) for header in headers if _field_for_header(header)}
        if not mapping:
            continue
        merged_mapping.update(mapping)
        for row in rows[1:]:
            candidate = {}
            for index, header in enumerate(headers):
                field = mapping.get(header)
                if not field or index >= len(row):
                    continue
                value = _normalize_value(field, row[index])
                if value not in (None, ''):
                    candidate[field] = value
            if candidate:
                candidate.setdefault('name', _candidate_name(candidate))
                candidate.setdefault('type', 'concert')
                candidate['source_sheet'] = sheet['name']
                all_candidates.append(candidate)
    return all_candidates, merged_mapping


def _extract_docx_text(path: Path):
    try:
        from docx import Document

        document = Document(str(path))
        lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    lines.append(' | '.join(cells))
        if lines:
            return '\n'.join(lines)
    except Exception:
        pass

    with zipfile.ZipFile(path) as archive:
        xml_content = archive.read('word/document.xml')
    root = ElementTree.fromstring(xml_content)
    texts = [node.text for node in root.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t') if node.text]
    return '\n'.join(texts)


def _extract_pdf_text(path: Path):
    try:
        import pdfplumber

        with pdfplumber.open(str(path)) as pdf:
            text = '\n'.join((page.extract_text() or '') for page in pdf.pages)
            if text.strip():
                return text
    except ImportError:
        pass

    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return '\n'.join((page.extract_text() or '') for page in reader.pages)
    except ImportError as exc:
        raise RuntimeError('缺少 PDF 解析依赖 pdfplumber 或 pypdf') from exc


def _extract_spreadsheet_sheets(path: Path):
    suffix = path.suffix.lower()
    if suffix == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as file:
            return [{'name': path.stem, 'rows': [row for row in csv.reader(file)]}]

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(str(path), data_only=True)
        sheets = []
        for sheet in workbook.worksheets:
            rows = [
                [cell for cell in row]
                for row in sheet.iter_rows(values_only=True)
                if any(cell not in (None, '') for cell in row)
            ]
            sheets.append({'name': sheet.title, 'rows': rows[:200]})
        return sheets
    except Exception:
        return _extract_xlsx_sheets_with_zip(path)


def _extract_xlsx_sheets_with_zip(path: Path):
    with zipfile.ZipFile(path) as archive:
        sheet_names = _xlsx_sheet_names(archive)
        shared_strings = _xlsx_shared_strings(archive)
        sheets = []
        for index, name in enumerate(sheet_names, start=1):
            sheet_path = f'xl/worksheets/sheet{index}.xml'
            if sheet_path not in archive.namelist():
                continue
            root = ElementTree.fromstring(archive.read(sheet_path))
            rows = []
            for row_node in root.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}row'):
                row_values = []
                for cell in row_node.findall('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c'):
                    row_values.append(_xlsx_cell_value(cell, shared_strings))
                rows.append(row_values)
            sheets.append({'name': name, 'rows': rows[:200]})
        return sheets


def _xlsx_sheet_names(archive):
    root = ElementTree.fromstring(archive.read('xl/workbook.xml'))
    names = [sheet.attrib.get('name', f'Sheet{index}') for index, sheet in enumerate(root.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet'), start=1)]
    return names or ['Sheet1']


def _xlsx_shared_strings(archive):
    if 'xl/sharedStrings.xml' not in archive.namelist():
        return []
    root = ElementTree.fromstring(archive.read('xl/sharedStrings.xml'))
    return [
        ''.join(node.text or '' for node in item.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t'))
        for item in root.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}si')
    ]


def _xlsx_cell_value(cell, shared_strings):
    cell_type = cell.attrib.get('t')
    if cell_type == 'inlineStr':
        text_nodes = cell.findall('.//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t')
        return ''.join(node.text or '' for node in text_nodes)
    value_node = cell.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v')
    if value_node is None or value_node.text is None:
        return ''
    if cell_type == 's':
        return shared_strings[int(value_node.text)]
    return _coerce_number(value_node.text)


def _field_for_header(header: str):
    normalized = _normalize_header(header)
    for field, aliases in FIELD_ALIASES.items():
        if normalized in {_normalize_header(alias) for alias in aliases}:
            return field
    return None


def _normalize_header(value: str):
    return re.sub(r'[\s_/-]+', '', str(value).strip().lower())


def _normalize_value(field: str, value):
    if value is None:
        return None
    if field in NUMBER_FIELDS:
        return _coerce_number(value)
    return str(value).strip()


def _coerce_number(value):
    if value in (None, ''):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else value
    text = re.sub(r'[,，¥￥\s]', '', str(value))
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return value
    return int(number) if number.is_integer() else number


def _candidate_name(candidate):
    return f"{candidate.get('artist_name') or '未命名项目'}-{candidate.get('city') or '待定城市'}"


def _line_title(line: str, fallback: str):
    title = re.split(r'[:：]', line, maxsplit=1)[0].strip()
    return title[:80] or fallback


def _unavailable_document_result(evidence, message: str):
    return {
        'file_kind': 'document',
        'parse_scope': 'single_project',
        'extracted_text': '',
        'candidate_facts': [{
            'title': f'{evidence.name} 解析候选事实',
            'content': message,
            'source': evidence.source or evidence.evidence_type,
        }],
        'candidate_assumptions': [{
            'title': '资料完整性假设',
            'content': '当前文件解析结果可能不完整，需人工录入或补充资料。',
            'confidence': 40,
        }],
        'candidate_risks': [{
            'title': '资料解析未核验风险',
            'level': 'medium',
            'mitigation': '由负责人确认解析字段、合同条款和关键数值后再进入最终判断。',
        }],
        'candidate_gates': [{
            'name': '资料人工核验',
            'status': 'pending',
            'required_evidence': evidence.name,
            'owner_group': 'B',
        }],
        'requires_human_verification': True,
        'error': message,
    }
