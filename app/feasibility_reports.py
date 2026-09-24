import io
import math

from docx import Document


REQUIRED_FINANCE_FIELDS = [
    "expected_attendance",
    "avg_ticket_price",
    "artist_fee",
    "venue_cost",
    "marketing_cost",
    "production_cost",
]


def build_feasibility_calculation(project, payload):
    values = {
        "expected_attendance": project.expected_attendance,
        "avg_ticket_price": project.avg_ticket_price,
        "artist_fee": project.artist_fee,
        "venue_cost": project.venue_cost,
        "marketing_cost": project.marketing_cost,
        "production_cost": project.production_cost,
    }
    missing = [field for field, value in values.items() if value is None]
    if missing:
        return {
            "formula_version": "feasibility-v1",
            "status": "pending_input",
            "missing_fields": missing,
            "scenarios": {},
            "total_cost": None,
            "breakeven_occupancy_rate": None,
        }

    capacity = values["expected_attendance"]
    ticket_price = values["avg_ticket_price"]
    tax_fee_rate = payload.tax_fee_rate if payload.tax_fee_rate is not None else 0.15
    sponsorship_income = payload.sponsorship_income or 0
    merchandise_income = payload.merchandise_income or 0
    total_cost = (
        values["artist_fee"]
        + values["venue_cost"]
        + values["marketing_cost"]
        + values["production_cost"]
    )
    full_net_ticket_revenue = capacity * ticket_price * (1 - tax_fee_rate)

    scenarios = {}
    for key, label, occupancy_rate in [
        ("conservative", "保守", 0.7),
        ("neutral", "中性", 0.8),
        ("optimistic", "乐观", 0.9),
    ]:
        attendance = round(capacity * occupancy_rate)
        gross_ticket_revenue = attendance * ticket_price
        net_ticket_revenue = round(gross_ticket_revenue * (1 - tax_fee_rate))
        total_income = net_ticket_revenue + sponsorship_income + merchandise_income
        net_profit = total_income - total_cost
        scenarios[key] = {
            "label": label,
            "occupancy_rate": occupancy_rate,
            "attendance": attendance,
            "avg_ticket_price": ticket_price,
            "gross_ticket_revenue": gross_ticket_revenue,
            "tax_fee_rate": tax_fee_rate,
            "net_ticket_revenue": net_ticket_revenue,
            "sponsorship_income": sponsorship_income,
            "merchandise_income": merchandise_income,
            "total_income": total_income,
            "total_cost": total_cost,
            "net_profit": net_profit,
            "roi": round(net_profit / total_cost, 4) if total_cost else None,
        }

    return {
        "formula_version": "feasibility-v1",
        "status": "calculated",
        "missing_fields": [],
        "capacity": capacity,
        "avg_ticket_price": ticket_price,
        "tax_fee_rate": tax_fee_rate,
        "sponsorship_income": sponsorship_income,
        "merchandise_income": merchandise_income,
        "total_cost": total_cost,
        "full_gross_ticket_revenue": capacity * ticket_price,
        "full_net_ticket_revenue": round(full_net_ticket_revenue),
        "breakeven_attendance": math.ceil(total_cost / (ticket_price * (1 - tax_fee_rate))) if ticket_price else None,
        "breakeven_occupancy_rate": round(total_cost / full_net_ticket_revenue, 4) if full_net_ticket_revenue else None,
        "scenarios": scenarios,
    }


def _money(value):
    return f"{int(value):,} 元" if value is not None else "待补齐"


def _percent(value):
    return f"{value * 100:.2f}%" if value is not None else "待补齐"


def build_feasibility_report_docx(project, calculation, facts, risks, assumptions, *, copy_mode="template_enriched"):
    document = Document()
    document.add_heading(f"{project.name} 可行性研究报告", level=0)
    document.add_paragraph(f"艺人：{project.artist_name or '待确认'}")
    document.add_paragraph(f"城市/场馆：{project.city or '待确认'} / {project.venue or '待确认'}")
    document.add_paragraph(f"档期：{project.schedule or '待确认'}")
    document.add_paragraph("说明：本报告中的硬性财务数据由系统参数计算生成，外部资料和人工录入内容均需保留核验。")

    sections = [
        "第一章 项目概述",
        "第二章 市场分析",
        "第三章 场地与档期比较",
        "第四章 投资预算",
        "第五章 定价策略与产出预估",
        "第六章 盈亏平衡与投资回报",
        "第七章 风险评估与应对",
        "第八章 结论与建议",
        "附录 数据来源",
    ]
    document.add_heading("目录", level=1)
    for section in sections:
        document.add_paragraph(section)

    document.add_heading("第一章 项目概述", level=1)
    document.add_paragraph(
        f"本项目拟围绕{project.artist_name or '目标艺人'}在{project.city or '目标城市'}开展{project.type or '演出'}项目，"
        f"计划场馆为{project.venue or '待确认场馆'}，预计可售规模按 {calculation.get('capacity', '待补齐')} 人测算。"
    )

    document.add_heading("第二章 市场分析", level=1)
    if facts:
        for fact in facts[:8]:
            document.add_paragraph(f"{fact.title}：{fact.content or '内容待补充'}（来源：{fact.source or '未标注'}；状态：{fact.status}）")
    else:
        document.add_paragraph("暂无已录入事实资料，建议补齐艺人热度、竞品项目、票务平台与社媒声量证据。")

    document.add_heading("第三章 场地与档期比较", level=1)
    document.add_paragraph(
        f"当前测算场馆为{project.venue or '待确认'}，档期为{project.schedule or '待确认'}。场馆容量、报批周期、周边交通与同档期竞品需在推进前完成复核。"
    )

    document.add_heading("第四章 投资预算", level=1)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "成本项"
    table.rows[0].cells[1].text = "金额"
    for label, value in [
        ("艺人成本", project.artist_fee),
        ("场馆成本", project.venue_cost),
        ("营销成本", project.marketing_cost),
        ("制作成本", project.production_cost),
        ("总成本", calculation.get("total_cost")),
    ]:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = _money(value)

    document.add_heading("第五章 定价策略与产出预估", level=1)
    document.add_paragraph(
        f"平均票价按 {_money(project.avg_ticket_price).replace(' 元', ' 元/人')} 测算；票务服务及税费率按 {_percent(calculation.get('tax_fee_rate'))} 扣减。"
    )
    scenario_table = document.add_table(rows=1, cols=7)
    headers = ["情景", "上座率", "人数", "毛票房", "净票房", "总收入", "净利润"]
    for index, header in enumerate(headers):
        scenario_table.rows[0].cells[index].text = header
    for scenario in calculation.get("scenarios", {}).values():
        row = scenario_table.add_row().cells
        row[0].text = scenario["label"]
        row[1].text = _percent(scenario["occupancy_rate"])
        row[2].text = f"{scenario['attendance']:,}"
        row[3].text = _money(scenario["gross_ticket_revenue"])
        row[4].text = _money(scenario["net_ticket_revenue"])
        row[5].text = _money(scenario["total_income"])
        row[6].text = _money(scenario["net_profit"])

    document.add_heading("第六章 盈亏平衡与投资回报", level=1)
    document.add_paragraph(
        f"盈亏平衡人数为 {calculation.get('breakeven_attendance'):,} 人，"
        f"盈亏平衡上座率为 {_percent(calculation.get('breakeven_occupancy_rate'))}。"
    )
    neutral = calculation.get("scenarios", {}).get("neutral") or {}
    document.add_paragraph(
        f"中性情境净利润为 {_money(neutral.get('net_profit'))}，ROI 为 {_percent(neutral.get('roi'))}。"
    )

    document.add_heading("第七章 风险评估与应对", level=1)
    if risks:
        for risk in risks[:8]:
            document.add_paragraph(f"{risk.title}：等级 {risk.level}；应对：{risk.mitigation or '待补充'}。")
    else:
        document.add_paragraph("暂无风险记录，建议补齐审批、艺人授权、场地安全、票务销售和资金合同风险。")

    document.add_heading("第八章 结论与建议", level=1)
    if neutral.get("net_profit", 0) >= 0:
        document.add_paragraph("按当前参数，中性情境具备正向收益空间，建议在完成政策审批、场地安全与艺人授权核验后进入下一阶段。")
    else:
        document.add_paragraph("按当前参数，中性情境未覆盖总成本，建议重新审视成本结构、票价策略、赞助收入与上座率假设。")
    if assumptions:
        document.add_paragraph("关键假设：")
        for assumption in assumptions[:6]:
            document.add_paragraph(f"{assumption.title}：{assumption.content or '待补充'}（置信度 {assumption.confidence}%）")

    document.add_heading("附录 数据来源", level=1)
    document.add_paragraph(f"报告生成模式：{copy_mode}")
    document.add_paragraph("财务测算口径：满座毛票房、扣除税费后的净票房、70%/80%/90% 三档上座率、总收入、净利润、ROI 与盈亏平衡上座率。")
    document.add_paragraph("外部采集、文档解析和 AI 生成建议均作为待核验证据，不直接替代人工决策。")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
